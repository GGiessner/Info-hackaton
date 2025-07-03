import os
import docx
import json
import openpyxl
import requests
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import pandas as pd

# === FONCTIONS UTILITAIRES ===

def lire_docx(path):
    doc = docx.Document(path)
    return "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])

def lire_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    contenu = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            ligne = [str(cell) for cell in row if cell is not None]
            if ligne:
                contenu.append(" | ".join(ligne))
    return "\n".join(contenu)

def extraire_texte_exemple(notes_path, devis_path, prop_path):
    notes = lire_docx(notes_path)
    prop = lire_docx(prop_path)
    return [notes, None, prop]

def extraire_texte_test(notes_path, devis_path, template_path):
    notes = lire_docx(notes_path)
    return [notes, None, None]

def generer_prompt(ex1, ex2, test_input, sections_demandées):
    sections = "\n".join(f"- {s}" for s in sections_demandées)
    return f"""
Tu es un assistant chargé de rédiger des propositions d'intervention professionnelles à partir de notes et de devis.

Voici deux exemples complets :

### Exemple 1 :
notes: {ex1[0]}
contrat: {ex1[2]}

### Exemple 2 :
notes: {ex2[0]}
contrat: {ex2[2]}

---

Voici maintenant l'étude de cas avec seulement les notes. C'est l'étude travaillée. C'EST CE SUR QUOI TU DOIS T'APPUYER POUR CE QUE JE TE DEMANDE PAR LA SUITE:
notes: {test_input[0]}

Génère uniquement les sections suivantes de la future proposition :
{sections}

Donne-moi la réponse au format JSON, avec un dictionnaire de la forme :
{{
  "Nom de la section 1": "Texte généré pour cette section",
  "Nom de la section 2": "..."
}}

Pour "Contexte":
Tu t'appuieras uniquement sur la note de l'étude travaillée.
Faire au moins 20 phrases.
Tu dois absolument inclure des données chiffrées, qui sont dans les notes
Indiquer la situation passée et actuelle du client.
Indiquer les besoins, les envies du client (expliqués en termes professionnels, dans le cadre d'une étude en conseil stratégique).
Je ne veux pas de réponse bateau. Je ne veux pas d'une réponse qui semble être applicable à toutes les entreprises, et qui est générale.
Je veux une réponse personnalisée!
"""

def interroger_mistral(prompt):
    try:
        response = requests.post(
            "http://localhost:11434/api/generate",
            json={"model": "mistral", "prompt": prompt, "stream": False}
        )
        response.raise_for_status()
        texte = response.json()["response"]
    except requests.RequestException as e:
        print(f"Erreur lors de la requête à Mistral : {e}")
        return {}

    try:
        return json.loads(texte)
    except json.JSONDecodeError:
        json_start = texte.find("{")
        json_end = texte.rfind("}") + 1
        try:
            return json.loads(texte[json_start:json_end])
        except:
            print("Erreur de parsing JSON. Contenu reçu :\n", texte)
            return {}

def trouver_descriptions_proches_mistral(texte_contexte, catalogue_textes):
    prompt = f"""
Voici un extrait de texte décrivant une situation client :

{texte_contexte}

Voici une liste de descriptions de projets :
{json.dumps(catalogue_textes, indent=2)}

Sélectionne les 9 descriptions les plus proches ou les plus utiles pour illustrer ce projet, dans le même format (liste JSON de chaînes).
"""
    try:
        response = requests.post(
            "http://localhost:11434/api/generate",
            json={"model": "mistral", "prompt": prompt, "stream": False}
        )
        response.raise_for_status()
        texte = response.json()["response"]
    except requests.RequestException as e:
        print(f"Erreur lors de la requête à Mistral pour les descriptions : {e}")
        return []

    try:
        descriptions = json.loads(texte)
        # Si la réponse contient des objets (dict), extraire uniquement les champs "description"
        if isinstance(descriptions, list) and descriptions and isinstance(descriptions[0], dict):
            descriptions = [d.get("description", "") for d in descriptions]
        # Supprimer doublons et chaînes vides
        descriptions_uniques = list(dict.fromkeys([d.strip() for d in descriptions if d.strip()]))
        # Limiter à 9 descriptions max (au cas où)
        return descriptions_uniques[:9]
    except json.JSONDecodeError:
        json_start = texte.find("[")
        json_end = texte.rfind("]") + 1
        try:
            descriptions = json.loads(texte[json_start:json_end])
            if isinstance(descriptions, list) and descriptions and isinstance(descriptions[0], dict):
                descriptions = [d.get("description", "") for d in descriptions]
            descriptions_uniques = list(dict.fromkeys([d.strip() for d in descriptions if d.strip()]))
            return descriptions_uniques[:9]
        except:
            print("Erreur de parsing JSON (descriptions). Contenu reçu :\n", texte)
            return []

# === CHEMINS DES FICHIERS ===

notes_ex1_path = "Gjoa MdP/Exemple 1/Exemple 1 - notes FC.docx"
devis_ex1_path = "Gjoa MdP/Exemple 1/20250631-Client-Sujet-Devis exemple 1.xlsx"
prop_ex1_path  = "Gjoa MdP/Exemple 1/20250631-Gjoa-Sujet-Proposition.docx"

notes_ex2_path = "Gjoa MdP/Exemple 2/Exemple 2 - notes FC.docx"
devis_ex2_path = "Gjoa MdP/Exemple 2/20250401-Devis-Client-Exemple 2.xlsx"
prop_ex2_path  = "Gjoa MdP/Exemple 2/20250402-Gjoa-Client-Exemple 2-Proposition.docx"

notes_3_path   = "Gjoa MdP/Test/Test - notes FC.docx"
devis_3_path   = "Gjoa MdP/Test/20250515-Devis-client-projet.xlsx"
Template       = "Template.docx"
catalogue_path = "Gjoa MdP/202209-Liste références et bilans projets - extract MdP.xlsx"

# === SECTIONS À GÉNÉRER ===

sections_voulues = {
    "Contexte": None
}

# === PIPELINE PRINCIPAL ===

if __name__ == "__main__":
    exemple1 = extraire_texte_exemple(notes_ex1_path, devis_ex1_path, prop_ex1_path)
    exemple2 = extraire_texte_exemple(notes_ex2_path, devis_ex2_path, prop_ex2_path)
    test_input = extraire_texte_test(notes_3_path, devis_3_path, Template)

    prompt = generer_prompt(exemple1, exemple2, test_input, sections_voulues.keys())
    resultats = interroger_mistral(prompt)

    if "Contexte" in resultats:
        df = pd.read_excel(catalogue_path)
        descriptions_catalogue = df["Description"].dropna().astype(str).tolist()
        descriptions_proches = trouver_descriptions_proches_mistral(resultats["Contexte"], descriptions_catalogue)
        resultats["Descriptions similaires issues du catalogue"] = descriptions_proches

    print(json.dumps(resultats, indent=2, ensure_ascii=False))